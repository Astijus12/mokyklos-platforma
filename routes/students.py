"""Mokinių valdymo puslapiai (CRUD - kurti, skaityti, atnaujinti, šalinti)."""
import csv
import io
import os
from datetime import datetime

from flask import Blueprint, render_template, request, redirect, url_for, flash, Response, current_app, send_from_directory
from flask_login import login_required, current_user
from werkzeug.utils import secure_filename
from routes.auth import personalas_required

from models import db, Mokinys, Klase

students_bp = Blueprint("students", __name__, url_prefix="/mokiniai")


def _parse_data(reiksme):
    if not reiksme:
        return None
    try:
        return datetime.strptime(reiksme, "%Y-%m-%d").date()
    except ValueError:
        return None


def _surinkti_duomenis(form):
    """Iš formos duomenų sukuria/atnaujina mokinio duomenis."""
    return {
        "vardas": form.get("vardas", "").strip(),
        "pavarde": form.get("pavarde", "").strip(),
        "gimimo_data": _parse_data(form.get("gimimo_data")),
        "email": form.get("email", "").strip() or None,
        "telefonas": form.get("telefonas", "").strip() or None,
        "adresas": form.get("adresas", "").strip() or None,
        "pastabos": form.get("pastabos", "").strip() or None,
        "motinos_vardas": form.get("motinos_vardas", "").strip() or None,
        "motinos_telefonas": form.get("motinos_telefonas", "").strip() or None,
        "motinos_email": form.get("motinos_email", "").strip() or None,
        "tevo_vardas": form.get("tevo_vardas", "").strip() or None,
        "tevo_telefonas": form.get("tevo_telefonas", "").strip() or None,
        "tevo_email": form.get("tevo_email", "").strip() or None,
        "mokyklos_autobusas": form.get("mokyklos_autobusas") == "on",
        "geroves_komisija": form.get("geroves_komisija") == "on",
        "geroves_komisija_data": _parse_data(form.get("geroves_komisija_data")) if form.get("geroves_komisija") == "on" else None,
        "geroves_komisija_pastabos": form.get("geroves_komisija_pastabos", "").strip() or None if form.get("geroves_komisija") == "on" else None,
    }


@students_bp.route("/")
@personalas_required
def saraso():
    paieska = request.args.get("paieska", "").strip()
    klases_id = request.args.get("klases_id", type=int)

    query = Mokinys.query
    if paieska:
        like = f"%{paieska}%"
        query = query.filter(
            (Mokinys.vardas.ilike(like)) | (Mokinys.pavarde.ilike(like))
        )
    if klases_id:
        query = query.filter_by(klases_id=klases_id)

    mokiniai = query.order_by(Mokinys.pavarde, Mokinys.vardas).all()
    klases = Klase.query.order_by(Klase.pavadinimas).all()

    return render_template(
        "students/list.html",
        mokiniai=mokiniai,
        klases=klases,
        paieska=paieska,
        pasirinkta_klase=klases_id,
    )


@students_bp.route("/eksportas")
@personalas_required
def eksportas_csv():
    """Mokinių sąrašo eksportas į CSV failą."""
    klases_id = request.args.get("klases_id", type=int)

    query = Mokinys.query
    if klases_id:
        query = query.filter_by(klases_id=klases_id)

    mokiniai = query.order_by(Mokinys.pavarde, Mokinys.vardas).all()

    output = io.StringIO()
    output.write("﻿")
    rasytojas = csv.writer(output, delimiter=";")
    rasytojas.writerow([
        "Vardas", "Pavardė", "Klasė", "Gimimo data",
        "El. paštas", "Telefonas", "Adresas",
        "Motinos vardas", "Motinos telefonas", "Motinos el. paštas",
        "Tėvo vardas", "Tėvo telefonas", "Tėvo el. paštas",
        "Mokyklos autobusas", "Gerovės komisija", "Gerovės kom. data",
    ])

    for m in mokiniai:
        rasytojas.writerow([
            m.vardas,
            m.pavarde,
            m.klase.pavadinimas if m.klase else "",
            m.gimimo_data.strftime("%Y-%m-%d") if m.gimimo_data else "",
            m.email or "",
            m.telefonas or "",
            m.adresas or "",
            m.motinos_vardas or "",
            m.motinos_telefonas or "",
            m.motinos_email or "",
            m.tevo_vardas or "",
            m.tevo_telefonas or "",
            m.tevo_email or "",
            "Taip" if m.mokyklos_autobusas else "Ne",
            "Taip" if m.geroves_komisija else "Ne",
            m.geroves_komisija_data.strftime("%Y-%m-%d") if m.geroves_komisija_data else "",
        ])

    failo_vardas = f"mokiniai_{datetime.now().strftime('%Y%m%d_%H%M')}.csv"
    return Response(
        output.getvalue().encode("utf-8"),
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={failo_vardas}"},
    )


@students_bp.route("/<int:mokinio_id>")
@personalas_required
def detales(mokinio_id):
    """Detalus mokinio puslapis."""
    mokinys = Mokinys.query.get_or_404(mokinio_id)
    return render_template("students/detail.html", mokinys=mokinys)


@students_bp.route("/<int:mokinio_id>/ataskaita")
@personalas_required
def ataskaita(mokinio_id):
    """Spausdinama mokinio ataskaita (PDF per naršyklės Print funkciją)."""
    mokinys = Mokinys.query.get_or_404(mokinio_id)
    return render_template("students/ataskaita.html", mokinys=mokinys, dabar=datetime.now())


@students_bp.route("/naujas", methods=["GET", "POST"])
@personalas_required
def naujas():
    if request.method == "POST":
        klases_id = request.form.get("klases_id", type=int)
        if not klases_id:
            flash("Pasirinkite klasę.", "danger")
        else:
            duomenys = _surinkti_duomenis(request.form)
            duomenys["klases_id"] = klases_id

            if not duomenys["vardas"] or not duomenys["pavarde"]:
                flash("Vardas ir pavardė yra privalomi.", "danger")
            else:
                mokinys = Mokinys(**duomenys)
                db.session.add(mokinys)
                db.session.commit()
                flash(f"Mokinys {mokinys.pilnas_vardas} pridėtas.", "success")
                return redirect(url_for("students.detales", mokinio_id=mokinys.id))

    klases = Klase.query.order_by(Klase.pavadinimas).all()
    return render_template(
        "students/form.html",
        mokinys=None,
        klases=klases,
        pradine_klase=request.args.get("klases_id", type=int),
    )


@students_bp.route("/<int:mokinio_id>/redaguoti", methods=["GET", "POST"])
@personalas_required
def redaguoti(mokinio_id):
    mokinys = Mokinys.query.get_or_404(mokinio_id)

    if request.method == "POST":
        duomenys = _surinkti_duomenis(request.form)
        for laukas, reiksme in duomenys.items():
            setattr(mokinys, laukas, reiksme)
        mokinys.klases_id = request.form.get("klases_id", type=int)
        db.session.commit()
        flash(f"Mokinio {mokinys.pilnas_vardas} duomenys atnaujinti.", "success")
        return redirect(url_for("students.detales", mokinio_id=mokinys.id))

    klases = Klase.query.order_by(Klase.pavadinimas).all()
    return render_template("students/form.html", mokinys=mokinys, klases=klases)


@students_bp.route("/<int:mokinio_id>/salinti", methods=["POST"])
@personalas_required
def salinti(mokinio_id):
    mokinys = Mokinys.query.get_or_404(mokinio_id)
    vardas = mokinys.pilnas_vardas
    db.session.delete(mokinys)
    db.session.commit()
    flash(f"Mokinys {vardas} pašalintas.", "info")
    return redirect(url_for("students.saraso"))


@students_bp.route("/grupinis-veiksmas", methods=["POST"])
@personalas_required
def grupinis_veiksmas():
    """Grupinis veiksmas su keliais mokiniais - eksportas arba šalinimas."""
    mokiniu_ids = [int(x) for x in request.form.getlist("mokiniu_ids") if x.isdigit()]
    veiksmas = request.form.get("veiksmas", "")

    if not mokiniu_ids:
        flash("Nepasirinktas nė vienas mokinys.", "warning")
        return redirect(url_for("students.saraso"))

    if veiksmas == "eksportuoti":
        return redirect(
            url_for("students.eksportas_pasirinktu", ids=",".join(str(i) for i in mokiniu_ids))
        )

    if veiksmas == "salinti":
        mokiniai = Mokinys.query.filter(Mokinys.id.in_(mokiniu_ids)).all()
        skaicius = len(mokiniai)
        for m in mokiniai:
            db.session.delete(m)
        db.session.commit()
        flash(f"Pašalinta {skaicius} mokinys(-iai).", "success")

    return redirect(url_for("students.saraso"))


@students_bp.route("/eksportas-pasirinktu")
@personalas_required
def eksportas_pasirinktu():
    """Eksportuoti pasirinktus mokinius į CSV."""
    ids_raw = request.args.get("ids", "")
    ids = [int(x) for x in ids_raw.split(",") if x.isdigit()]

    if not ids:
        return redirect(url_for("students.saraso"))

    mokiniai = Mokinys.query.filter(Mokinys.id.in_(ids)).order_by(
        Mokinys.pavarde, Mokinys.vardas
    ).all()

    output = io.StringIO()
    output.write("﻿")
    rasytojas = csv.writer(output, delimiter=";")
    rasytojas.writerow([
        "Vardas", "Pavardė", "Klasė", "Gimimo data",
        "El. paštas", "Telefonas", "Adresas",
        "Motinos vardas", "Motinos telefonas", "Motinos el. paštas",
        "Tėvo vardas", "Tėvo telefonas", "Tėvo el. paštas",
        "Mokyklos autobusas", "Gerovės komisija",
    ])

    for m in mokiniai:
        rasytojas.writerow([
            m.vardas, m.pavarde,
            m.klase.pavadinimas if m.klase else "",
            m.gimimo_data.strftime("%Y-%m-%d") if m.gimimo_data else "",
            m.email or "", m.telefonas or "", m.adresas or "",
            m.motinos_vardas or "", m.motinos_telefonas or "", m.motinos_email or "",
            m.tevo_vardas or "", m.tevo_telefonas or "", m.tevo_email or "",
            "Taip" if m.mokyklos_autobusas else "Ne",
            "Taip" if m.geroves_komisija else "Ne",
        ])

    return Response(
        output.getvalue().encode("utf-8"),
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename=pasirinkti_mokiniai_{datetime.now().strftime('%Y%m%d_%H%M')}.csv"},
    )


@students_bp.route("/eksportas-xlsx")
@personalas_required
def eksportas_xlsx():
    """Mokinių sąrašo eksportas į Excel formato failą su formatavimu."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    klases_id = request.args.get("klases_id", type=int)
    query = Mokinys.query
    if klases_id:
        query = query.filter_by(klases_id=klases_id)
    mokiniai = query.order_by(Mokinys.pavarde, Mokinys.vardas).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Mokiniai"

    antraste_font = Font(bold=True, color="FFFFFF", size=11, name="Calibri")
    antraste_fill = PatternFill("solid", fgColor="881337")
    centruota = Alignment(horizontal="center", vertical="center", wrap_text=True)
    kraštine = Border(
        left=Side(border_style="thin", color="D0D0D0"),
        right=Side(border_style="thin", color="D0D0D0"),
        top=Side(border_style="thin", color="D0D0D0"),
        bottom=Side(border_style="thin", color="D0D0D0"),
    )

    antrastes = [
        "Vardas", "Pavardė", "Klasė", "Gimimo data",
        "El. paštas", "Telefonas", "Adresas",
        "Motinos vardas", "Motinos tel.",
        "Tėvo vardas", "Tėvo tel.",
        "Autobusas", "VGK",
    ]

    for stulpelio_nr, tekstas in enumerate(antrastes, 1):
        c = ws.cell(row=1, column=stulpelio_nr, value=tekstas)
        c.font = antraste_font
        c.fill = antraste_fill
        c.alignment = centruota
        c.border = kraštine

    ws.row_dimensions[1].height = 26

    for eil_nr, m in enumerate(mokiniai, 2):
        eilute = [
            m.vardas,
            m.pavarde,
            m.klase.pavadinimas if m.klase else "",
            m.gimimo_data.strftime("%Y-%m-%d") if m.gimimo_data else "",
            m.email or "",
            m.telefonas or "",
            m.adresas or "",
            m.motinos_vardas or "",
            m.motinos_telefonas or "",
            m.tevo_vardas or "",
            m.tevo_telefonas or "",
            "Taip" if m.mokyklos_autobusas else "Ne",
            "Taip" if m.geroves_komisija else "Ne",
        ]
        for stulpelio_nr, reiksme in enumerate(eilute, 1):
            c = ws.cell(row=eil_nr, column=stulpelio_nr, value=reiksme)
            c.border = kraštine
            if stulpelio_nr in (12, 13):
                c.alignment = Alignment(horizontal="center")

    stulpeliu_pločiai = [14, 16, 8, 14, 26, 16, 30, 22, 16, 22, 16, 12, 8]
    for i, pločis in enumerate(stulpeliu_pločiai, 1):
        ws.column_dimensions[chr(64 + i)].width = pločis

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{chr(64 + len(antrastes))}{len(mokiniai) + 1}"

    buferis = io.BytesIO()
    wb.save(buferis)
    buferis.seek(0)

    failo_vardas = f"mokiniai_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return Response(
        buferis.read(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={failo_vardas}"},
    )


def _leidziamas_nuotraukos_failas(vardas):
    leidziami = {"png", "jpg", "jpeg", "webp"}
    return "." in vardas and vardas.rsplit(".", 1)[1].lower() in leidziami


@students_bp.route("/<int:mokinio_id>/nuotrauka", methods=["POST"])
@personalas_required
def ikelti_nuotrauka(mokinio_id):
    """Įkelia mokinio nuotrauką."""
    mokinys = Mokinys.query.get_or_404(mokinio_id)
    failas = request.files.get("nuotrauka")

    if not failas or not failas.filename:
        flash("Pasirinkite failą.", "danger")
        return redirect(url_for("students.detales", mokinio_id=mokinys.id))

    if not _leidziamas_nuotraukos_failas(failas.filename):
        flash("Neleistinas failo formatas. Naudokite PNG, JPG arba WebP.", "danger")
        return redirect(url_for("students.detales", mokinio_id=mokinys.id))

    saugus_vardas = secure_filename(failas.filename)
    plėtinys = saugus_vardas.rsplit(".", 1)[1].lower()
    naujas_vardas = f"mokinys_{mokinys.id}_{datetime.now().strftime('%Y%m%d%H%M%S')}.{plėtinys}"

    nuotrauku_dir = os.path.join(current_app.config["UPLOAD_FOLDER"], "mokiniai")
    os.makedirs(nuotrauku_dir, exist_ok=True)
    kelias = os.path.join(nuotrauku_dir, naujas_vardas)
    failas.save(kelias)

    if mokinys.nuotraukos_failas:
        senas_kelias = os.path.join(current_app.config["UPLOAD_FOLDER"], mokinys.nuotraukos_failas)
        if os.path.exists(senas_kelias):
            try:
                os.remove(senas_kelias)
            except OSError:
                pass

    mokinys.nuotraukos_failas = f"mokiniai/{naujas_vardas}"
    db.session.commit()

    flash(f"Mokinio {mokinys.pilnas_vardas} nuotrauka atnaujinta.", "success")
    return redirect(url_for("students.detales", mokinio_id=mokinys.id))


@students_bp.route("/<int:mokinio_id>/nuotrauka", methods=["DELETE", "POST"], endpoint="salinti_nuotrauka_endpointas")
@personalas_required
def salinti_nuotrauka(mokinio_id):
    if request.form.get("_veiksmas") != "salinti":
        return redirect(url_for("students.detales", mokinio_id=mokinio_id))

    mokinys = Mokinys.query.get_or_404(mokinio_id)
    if mokinys.nuotraukos_failas:
        kelias = os.path.join(current_app.config["UPLOAD_FOLDER"], mokinys.nuotraukos_failas)
        if os.path.exists(kelias):
            try:
                os.remove(kelias)
            except OSError:
                pass
        mokinys.nuotraukos_failas = None
        db.session.commit()
        flash("Nuotrauka pašalinta.", "info")
    return redirect(url_for("students.detales", mokinio_id=mokinio_id))


@students_bp.route("/nuotrauka/<path:kelias>")
@personalas_required
def nuotraukos_failas(kelias):
    """Saugiai serveris nuotraukas iš uploads katalogo."""
    return send_from_directory(current_app.config["UPLOAD_FOLDER"], kelias)
