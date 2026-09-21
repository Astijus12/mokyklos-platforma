"""Užpildo Baigiamosios praktikos individualios užduoties dokumentą."""
import sys
from docx import Document
from docx.shared import Pt

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

ORIGINAL = r"C:\Users\astij\mokyklos-platforma\baigiamosios_uzduotis.docx"
NAUJAS = r"C:\Users\astij\OneDrive\Desktop\PI23SN_Grineviciaus_Baigiamosios_praktikos_UZPILDYTA.docx"


def prideti_run(p, tekstas, bold=False, underline=False):
    run = p.add_run(tekstas)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    if bold:
        run.bold = True
    if underline:
        run.underline = True
    return run


doc = Document(ORIGINAL)

# [8] Studentas + Grupė (jau užpildytas šablone: "Astijus Grinevičius, PI23SN Grupė")
p8 = doc.paragraphs[8]
for r in p8.runs:
    r.text = ""
prideti_run(p8, "Studentas(ė) ", bold=False)
prideti_run(p8, "Astijus Grinevičius", bold=True)
prideti_run(p8, ",\t\t\tGrupė ", bold=False)
prideti_run(p8, "PI23SN", bold=True)

# [10] Praktikos pavadinimas - PABRAUKTI "baigiamoji"
p10 = doc.paragraphs[10]
for r in p10.runs:
    r.text = ""
prideti_run(p10, "Praktikos pavadinimas: profesinė / ", bold=False)
prideti_run(p10, "baigiamoji", bold=True, underline=True)
prideti_run(p10, ".", bold=False)
prideti_run(p10, "(tinkamą pabraukti)", bold=False)

# [11] TIKSLAS
p11 = doc.paragraphs[11]
for r in p11.runs:
    r.text = ""
prideti_run(p11, "Tikslas: ", bold=False)
prideti_run(
    p11,
    "Patobulinti ir išplėsti profesinės praktikos metu sukurtą Kupiškio r. "
    "Subačiaus gimnazijos Noriūnų Jono Černiaus skyriaus mokyklos valdymo "
    "informacinę sistemą, pereinant prie produkcinio lygio architektūros, "
    "realizuojant tėvų prieigos portalą, sustiprinant saugumo mechanizmus "
    "bei atliekant sistemos naudojimo tyrimą realioje mokyklos aplinkoje.",
    bold=False
)

# UŽDAVINIAI (paragrafai 13-16)
uzdaviniai = [
    "Modernizuoti Sistemos technologinę architektūrą: atlikti duomenų bazės "
    "migraciją iš SQLite į PostgreSQL, įdiegti Docker konteinerizaciją, "
    "sukurti automatinę duomenų atsarginių kopijų sistemą bei paruošti CI/CD "
    "procesą, siekiant produkcinio lygio patikimumo ir plečiamumo.",

    "Suprojektuoti ir realizuoti tėvų prieigos portalą, leidžiantį tėvams "
    "saugiai pasiekti savo vaiko duomenis, oficialius mokyklos pranešimus, "
    "komunikuoti su klasės vadovu bei matyti vaiko gerovės komisijos "
    "informaciją.",

    "Sustiprinti Sistemos saugumo mechanizmus įdiegus dviejų veiksnių "
    "autentikaciją (2FA) administratorių paskyroms bei realizuoti REST API "
    "sąsają, leidžiančią mobiliesiems klientams pasiekti pagrindines "
    "Sistemos funkcijas.",

    "Parengti automatinę Sistemos testavimo sistemą (vienetiniai, "
    "integraciniai testai) bei atlikti realaus naudojimo tyrimą Subačiaus "
    "gimnazijos Noriūnų Jono Černiaus skyriuje - įvertinti vartotojų "
    "patirtį, surinkti mokytojų ir tėvų atsiliepimus bei parengti "
    "rekomendacijas tolimesniam Sistemos vystymui.",
]

for i, uzdavinio_tekstas in enumerate(uzdaviniai):
    p = doc.paragraphs[13 + i]
    for r in p.runs:
        r.text = ""
    prideti_run(p, uzdavinio_tekstas, bold=False)

# [20] Priimančios organizacijos praktikos vadovas
p20 = doc.paragraphs[20]
for r in p20.runs:
    r.text = ""
prideti_run(p20, "Priimančios organizacijos praktikos vadovas ", bold=False)
prideti_run(p20, "Rita Grinevičienė, Noriūnų Jono Černiaus skyriaus vedėja", bold=True)

# [22] EIF praktikos vadovas
p22 = doc.paragraphs[22]
for r in p22.runs:
    r.text = ""
prideti_run(p22, "EIF praktikos vadovas ", bold=False)
prideti_run(p22, "doc. dr. Romanas Tumasonis, lektorius", bold=True)

# [24] Data
p24 = doc.paragraphs[24]
for r in p24.runs:
    r.text = ""
prideti_run(p24, "2026 m. rugsėjo 1 d.", bold=False)

doc.save(NAUJAS)
print(f"✓ Užpildytas dokumentas išsaugotas:")
print(f"  {NAUJAS}")
print()
print("PAGRINDINIAI DUOMENYS:")
print("  Studentas:        Astijus Grinevičius")
print("  Grupė:            PI23SN")
print("  Praktikos tipas:  baigiamoji")
print("  Praktikos vieta:  Kupiškio r. Subačiaus gimnazija")
print("                    (Noriūnų Jono Černiaus skyrius)")
print("  Vadovė mokykloje: Rita Grinevičienė (skyriaus vedėja)")
print("  EIF vadovas:      doc. dr. Romanas Tumasonis")
print("  Data:             2026 m. rugsėjo 1 d.")
print()
print("4 UŽDAVINIAI:")
print("  1. Technologinės architektūros modernizavimas (PostgreSQL, Docker, backup)")
print("  2. Tėvų prieigos portalo kūrimas")
print("  3. Saugumo stiprinimas (2FA) + REST API mobiliesiems")
print("  4. Testavimo sistema + realaus naudojimo tyrimas")
