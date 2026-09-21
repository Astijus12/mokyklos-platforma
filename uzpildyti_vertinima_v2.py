"""Užpildo Praktikos vadovo vertinimo lapą (v2 - su pataisytomis pareigomis).

Rita Grinevičienė - SKYRIAUS VEDĖJA (ne direktoriaus pavaduotoja).
"""
import sys
from docx import Document
from docx.shared import Pt

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

ORIGINAL = r"C:\Users\astij\mokyklos-platforma\vertinimo_kopija.docx"
UZPILDYTAS = r"C:\Users\astij\mokyklos-platforma\PI23SN_profesines_vertinimas_v2.docx"


def pakeisti_pastraipa(p, naujas_tekstas, bold=False):
    """Pakeičia paragrafo turinį."""
    for r in p.runs:
        r.text = ""

    if p.runs:
        p.runs[0].text = naujas_tekstas
        p.runs[0].font.name = "Times New Roman"
        p.runs[0].font.size = Pt(12)
        if bold:
            p.runs[0].font.bold = True
    else:
        run = p.add_run(naujas_tekstas)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        if bold:
            run.font.bold = True


doc = Document(ORIGINAL)

# Studentas
pakeisti_pastraipa(
    doc.paragraphs[4],
    "Studentas(ė)\tAstijus Grinevičius"
)

# Praktikos atlikimo vieta (1)
pakeisti_pastraipa(
    doc.paragraphs[10],
    "Praktikos atlikimo vieta\tKupiškio r. Subačiaus gimnazija"
)

# Praktikos atlikimo vieta (2)
pakeisti_pastraipa(
    doc.paragraphs[11],
    "\t\t\t(Noriūnų Jono Černiaus skyrius), Melioratorių g. 5, Noriūnai k., Kupiškio r."
)

# 1. Požiūris į darbą
pakeisti_pastraipa(
    doc.paragraphs[21],
    "1. Požiūris į darbą (domėjimasis darbu, iniciatyvumas, pareigingumas, "
    "tvarkingumas, drausmė ir kt.): Studento požiūris į praktiką buvo pavyzdinis. "
    "Astijus aktyviai domėjosi mokyklos administraciniais procesais, savarankiškai "
    "iniciavo susitikimus su pedagoginiu personalu reikalavimams aiškintis. "
    "Pasižymėjo dideliu pareigingumu, atvyko laiku, visas užduotis atliko "
    "kruopščiai ir terminuotai. Bendraudamas su mokyklos darbuotojais buvo "
    "mandagus, dėmesingas ir profesionalus. Demonstravo aukštą motyvaciją bei "
    "stiprų norą sukurti realiai naudingą produktą mokyklos bendruomenei."
)

# 2. Teorinio parengimo trūkumai
pakeisti_pastraipa(
    doc.paragraphs[23],
    "2. Teorinio parengimo trūkumai, išryškėję praktikos metu: Esminių teorinio "
    "parengimo trūkumų praktikos metu nepastebėta. Studento technologinės žinios "
    "(Python, Flask, SQLite, HTML, CSS, JavaScript) atitiko keliamus reikalavimus. "
)

pakeisti_pastraipa(
    doc.paragraphs[24],
    "Nedidelių sunkumų buvo susijusių su projekto įdiegimu į debesijos paslaugą "
    "(Render.com), tačiau studentas savarankiškai išsiaiškino diegimo niuansus "
)

pakeisti_pastraipa(
    doc.paragraphs[25],
    "ir sėkmingai pristatė veikiantį produktą. Rekomenduotina ateityje gilinti "
)

pakeisti_pastraipa(
    doc.paragraphs[26],
    "žinias apie produkcinių sistemų diegimą bei duomenų bazių migracijų valdymą."
)

pakeisti_pastraipa(doc.paragraphs[27], "")
pakeisti_pastraipa(doc.paragraphs[28], "")

# 3. Atliktų užduočių kokybė
pakeisti_pastraipa(
    doc.paragraphs[30],
    "3. Atliktų užduočių kokybė, savarankiškumas: Praktikos metu studentas "
    "savarankiškai suprojektavo ir realizavo internetinę mokyklos valdymo "
    "informacinę sistemą, skirtą Subačiaus gimnazijos Noriūnų Jono Černiaus "
)

pakeisti_pastraipa(
    doc.paragraphs[31],
    "skyriaus mokytojams ir administracijai. Sistemoje įgyvendintos visos "
    "sutartos funkcijos: vartotojų autentikacija su rolėmis paremta prieiga, "
    "mokinių, klasių, tėvų duomenų valdymas, vaiko gerovės komisijos "
    "informacijos saugojimas, "
)

pakeisti_pastraipa(
    doc.paragraphs[32],
    "diplomų bei kvalifikacijos kėlimo dokumentų įkėlimas, vidiniai skelbimai, "
    "globali paieška bei CSV eksportas. Sistema sėkmingai įdiegta į debesijos "
    "platformą Render.com ir yra pasiekiama internetu. Visi darbai atlikti "
    "savarankiškai, kokybiškai, laikantis šiuolaikinių programinės įrangos "
    "kūrimo gerosios praktikos principų."
)

# 4. Praktinio darbo įvertinimas balais
pakeisti_pastraipa(
    doc.paragraphs[35],
    "4. Praktinio darbo įvertinimas balais\t10 (puikiai)"
)

# Priimančios organizacijos praktikos vadovas
pakeisti_pastraipa(
    doc.paragraphs[38],
    "Priimančios organizacijos praktikos vadovas\tRita Grinevičienė"
)

# PATAISYTA: skyriaus vedėja
pakeisti_pastraipa(
    doc.paragraphs[39],
    "\t\t\t\t(Noriūnų Jono Černiaus skyriaus vedėja)"
)

# Data
pakeisti_pastraipa(
    doc.paragraphs[44],
    "2026 m.\tbirželio\tmėn.\t19\td."
)

doc.save(UZPILDYTAS)
print(f"✓ Pataisytas vertinimo lapas išsaugotas:")
print(f"  {UZPILDYTAS}")
print()
print("PAKEITIMAS:")
print("  PAREIGOS: 'direktoriaus pavaduotoja ugdymui' → 'Noriūnų Jono Černiaus skyriaus vedėja'")
