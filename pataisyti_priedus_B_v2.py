"""Pataiso priedų struktūrą (Variantas B) - V2.

Pataisymas: pašalina pastraipas su pattern "X priedas." (mažosiomis raidėmis),
kurios yra po PRIEDAI antraštės.
"""
import sys
import re
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

ORIGINAL = r"C:\Users\astij\mokyklos-platforma\ataskaita_su_2priedu.docx"
NAUJAS = r"C:\Users\astij\mokyklos-platforma\ataskaita_su_2priedu_FIXED.docx"

doc = Document(ORIGINAL)

# Rasti PRIEDAI antraštę
priedai_idx = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().upper() == 'PRIEDAI' and p.style.name == 'Heading 1':
        priedai_idx = i
        break

print(f"PRIEDAI sekcija: [{priedai_idx}]")
print()

# Surasti ir ištrinti "X priedas." pastraipas (mažosiomis raidėmis)
# Pattern: prasideda "1 priedas." arba "2 priedas." ir t.t.
salinamos = []
for i in range(priedai_idx + 1, min(priedai_idx + 20, len(doc.paragraphs))):
    text = doc.paragraphs[i].text.strip()
    if re.match(r'^\d+\s+priedas\.', text):
        salinamos.append(i)
        print(f"  Sąraše: [{i}] {text[:80]}")

print()
print(f"Šalinamos {len(salinamos)} pastraipos...")

# Atvirkščia tvarka, kad indeksai nepasikeistų
for i in sorted(salinamos, reverse=True):
    p_element = doc.paragraphs[i]._element
    p_element.getparent().remove(p_element)

doc.save(NAUJAS)
print()
print(f"✓ Pataisyta ataskaita išsaugota:")
print(f"  {NAUJAS}")
